"""One-time Google setup through gws: create the resume template Doc and the tracker Sheet.

    python scripts/google_setup.py              # create whatever is missing, save IDs to .env
    python scripts/google_setup.py --template   # rebuild the template only (e.g. after editing the base resume)

The template is generated from resume/base_resume.md (DOCX with {{PLACEHOLDERS}}, uploaded
and converted to a Google Doc). Restyle it freely in Google Docs — fonts, colours, spacing —
as long as every {{...}} placeholder stays intact; all future PDFs follow its design.
"""

from __future__ import annotations

import argparse
import re
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from common import DATA_DIR, ROOT, base_resume_text
from gws import check_auth, gws
from resume_model import parse
from resume_template import template_markdown

ACCENT = "1F3A5F"
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*)")
ENV_PATH = ROOT / ".env"
GDOC_MIME = "application/vnd.google-apps.document"


# ---------------------------------------------------------------- DOCX (template source only)


def _runs(text: str) -> list[tuple[str, bool, bool]]:
    out = []
    for part in INLINE.split(text):
        if part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], True, False))
        elif part.startswith("*") and part.endswith("*") and len(part) > 1:
            out.append((part[1:-1], False, True))
        elif part:
            out.append((part, False, False))
    return out


def _blocks(md: str):
    r = parse(md)
    yield "name", r.name
    yield "contact", r.contact
    for section, lines in r.sections.items():
        yield "h2", section
        for line in lines:
            s = line.strip()
            if not s:
                continue
            if s.startswith("### "):
                yield "h3", s[4:]
            elif re.fullmatch(r"[*_][^*_].*[*_]", s) and not s.startswith("**"):
                yield "meta", s.strip("*_")
            elif re.match(r"[-*]\s+", s):
                yield "bullet", re.sub(r"^[-*]\s+", "", s)
            else:
                yield "para", s


def _bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    bdr, bottom = OxmlElement("w:pBdr"), OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", ACCENT)):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    p_pr.append(bdr)


MD_LINK = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)")  # [GitHub](https://github.com/...)


def _add_hyperlink(paragraph, text: str, url: str, size: float) -> None:
    """A clickable word (e.g. "GitHub") pointing at url."""
    rel_id = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                                      is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run, props = OxmlElement("w:r"), OxmlElement("w:rPr")
    for tag, attrs in (("w:color", {"w:val": ACCENT}), ("w:u", {"w:val": "single"}),
                       ("w:sz", {"w:val": str(int(size * 2))})):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        props.append(el)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(props)
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def build_docx(md: str, path: Path) -> None:
    doc = Document()
    for s in doc.sections:
        s.page_height, s.page_width = Pt(842), Pt(595)  # A4
        s.top_margin = s.bottom_margin = Pt(28)
        s.left_margin = s.right_margin = Pt(46)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)

    def add(text, *, size=None, bold=False, italic=False, color=None, align=None, style=None, before=0, after=2):
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        if align is not None:
            p.alignment = align
        for t, b, i in _runs(text):
            run = p.add_run(t)
            run.bold, run.italic = bold or b, italic or i
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
        return p

    for kind, text in _blocks(md):
        if kind == "name":
            add(text, size=20, bold=True, color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
        elif kind == "contact":
            p = add("", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
            pos = 0
            for m in MD_LINK.finditer(text):
                if m.start() > pos:
                    p.add_run(text[pos:m.start()]).font.size = Pt(10)
                _add_hyperlink(p, m.group(1), m.group(2), 10)
                pos = m.end()
            if pos < len(text):
                p.add_run(text[pos:]).font.size = Pt(10)
        elif kind == "h2":
            _bottom_border(add(text.upper(), size=11.5, bold=True, color=ACCENT, before=6, after=2))
        elif kind == "h3":
            add(text, bold=True, before=3, after=0)
        elif kind == "meta":
            add(text, italic=True, size=9.5, color="555555", after=1)
        elif kind == "bullet":
            add(text, style="List Bullet", after=0)
        else:
            add(text)
    doc.save(path)


# ---------------------------------------------------------------- Google resources


def set_env(key: str, value: str) -> None:
    text = ENV_PATH.read_text(encoding="utf-8") if ENV_PATH.exists() else ""
    line = f"{key}={value}"
    if re.search(rf"^{key}=.*$", text, re.M):
        text = re.sub(rf"^{key}=.*$", line, text, flags=re.M)
    else:
        text = text.rstrip("\n") + ("\n" if text else "") + line + "\n"
    ENV_PATH.write_text(text, encoding="utf-8")


def create_template() -> str:
    # gws only uploads files from inside the project directory
    with tempfile.TemporaryDirectory(dir=ensure_dir(DATA_DIR / "tmp")) as tmp:
        src = Path(tmp) / "resume_template.docx"
        build_docx(template_markdown(base_resume_text()), src)
        doc = gws("drive", "files", "create", body={"name": "Srikanth Valluri Resume - TEMPLATE (do not delete)",
                                                    "mimeType": GDOC_MIME}, upload=src, params={"fields": "id"})
    return doc["id"]


def ensure_dir(p: Path) -> Path:
    p.mkdir(parents=True, exist_ok=True)
    return p


def create_sheet() -> str:
    sh = gws("sheets", "spreadsheets", "create", body={
        "properties": {"title": "Job Applications Tracker"},
        "sheets": [{"properties": {"title": "Jobs", "gridProperties": {"frozenRowCount": 1}}}],
    })
    return sh["spreadsheetId"]


def _rgb(hex_color: str) -> dict:
    h = hex_color.lstrip("#")
    return {k: int(h[i:i + 2], 16) / 255 for k, i in (("red", 0), ("green", 2), ("blue", 4))}


def format_sheet(sid: str, rows: int = 5000) -> None:
    """Header, Status dropdown + colours, date check on Applied On, column widths. Safe to re-run."""
    from write_sheet import COLUMNS, STATUS_COL, STATUSES

    meta = gws("sheets", "spreadsheets", "get", params={
        "spreadsheetId": sid, "fields": "sheets(properties(sheetId,gridProperties),conditionalFormats)"})
    sheet = meta["sheets"][0]
    gid, grid = sheet["properties"]["sheetId"], sheet["properties"]["gridProperties"]
    gws("sheets", "spreadsheets", "values", "update", body={"values": [COLUMNS]},
        params={"spreadsheetId": sid, "range": "A1", "valueInputOption": "RAW"})

    def col_range(col: int) -> dict:
        return {"sheetId": gid, "startRowIndex": 1, "endRowIndex": rows, "startColumnIndex": col,
                "endColumnIndex": col + 1}

    # Start clean so re-runs (e.g. after columns change) don't leave rules on the wrong column.
    requests = [{"deleteConditionalFormatRule": {"sheetId": gid, "index": 0}}
                for _ in sheet.get("conditionalFormats", [])]
    requests.append({"setDataValidation": {"range": {"sheetId": gid, "startRowIndex": 1, "endRowIndex": rows}}})
    if grid.get("rowCount", 0) < rows:
        requests.append({"appendDimension": {"sheetId": gid, "dimension": "ROWS",
                                             "length": rows - grid.get("rowCount", 0)}})
    requests += [
        {"updateSheetProperties": {"properties": {"sheetId": gid, "gridProperties": {"frozenRowCount": 1}},
                                   "fields": "gridProperties.frozenRowCount"}},
        {"repeatCell": {"range": {"sheetId": gid, "startRowIndex": 0, "endRowIndex": 1},
                        "cell": {"userEnteredFormat": {"backgroundColor": _rgb("#1F3A5F"),
                                                       "textFormat": {"bold": True, "foregroundColor": _rgb("#FFFFFF")}}},
                        "fields": "userEnteredFormat(textFormat,backgroundColor)"}},
        {"setDataValidation": {"range": col_range(STATUS_COL), "rule": {
            "condition": {"type": "ONE_OF_LIST", "values": [{"userEnteredValue": s} for s in STATUSES]},
            "strict": True, "showCustomUi": True}}},
        {"setDataValidation": {"range": col_range(COLUMNS.index("Applied On")), "rule": {
            "condition": {"type": "DATE_IS_VALID"}, "strict": False,
            "inputMessage": "Date you applied, e.g. 2026-10-15"}}},
    ]
    for i, (status, color) in enumerate(STATUSES.items()):
        requests.append({"addConditionalFormatRule": {"index": i, "rule": {
            "ranges": [col_range(STATUS_COL)],
            "booleanRule": {"condition": {"type": "TEXT_EQ", "values": [{"userEnteredValue": status}]},
                            "format": {"backgroundColor": _rgb(color)}}}}})
    widths = {"Job Title": 230, "Company": 150, "Job Link": 160, "ATS Match": 220, "Fit Reason": 320,
              "Top Missing Skills": 260,
              "Talking Points": 380, "Resume (Drive)": 160, "Desktop PDF": 160, "Status": 110, "Notes": 240}
    for name, px in widths.items():
        c = COLUMNS.index(name)
        requests.append({"updateDimensionProperties": {
            "range": {"sheetId": gid, "dimension": "COLUMNS", "startIndex": c, "endIndex": c + 1},
            "properties": {"pixelSize": px}, "fields": "pixelSize"}})
    gws("sheets", "spreadsheets", "batchUpdate", params={"spreadsheetId": sid}, body={"requests": requests})


def main() -> int:
    import os

    from gdrive import ensure_folder, move_into

    ap = argparse.ArgumentParser(description="Create the Drive folder, Docs template and tracker Sheet via gws")
    ap.add_argument("--template", action="store_true", help="(re)build the template even if one exists")
    args = ap.parse_args()
    check_auth()

    root = os.getenv("GOOGLE_DRIVE_FOLDER_ID")
    if not root:
        root = ensure_folder("Job Applications")
        set_env("GOOGLE_DRIVE_FOLDER_ID", root)
        os.environ["GOOGLE_DRIVE_FOLDER_ID"] = root
    print(f"Drive folder: https://drive.google.com/drive/folders/{root}")

    tid = os.getenv("GOOGLE_RESUME_TEMPLATE_ID")
    if args.template or not tid:
        tid = create_template()
        set_env("GOOGLE_RESUME_TEMPLATE_ID", tid)
    sid = os.getenv("GOOGLE_SHEET_ID")
    if not sid:
        sid = create_sheet()
        set_env("GOOGLE_SHEET_ID", sid)
    for fid in (tid, sid):
        move_into(fid, root)  # keep everything together under "Job Applications"
    format_sheet(sid)
    print(f"Template:     https://docs.google.com/document/d/{tid}/edit")
    print(f"Tracker:      https://docs.google.com/spreadsheets/d/{sid}/edit")
    print("IDs saved to .env")
    return 0


if __name__ == "__main__":
    sys.exit(main())
